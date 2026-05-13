import csv
import ipaddress
import json
import logging
import os
import socket
import sqlite3
import ssl
import time
from collections import deque
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import netflow
from netflow.ipfix import IPFIXTemplateNotRecognized
from netflow.utils import UnknownExportVersion
from netflow.v9 import V9TemplateNotRecognized
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None

try:
    from librouteros import connect as routeros_connect
except ImportError:
    routeros_connect = None


if load_dotenv:
    load_dotenv()


def load_env_file_fallback(path: Path = Path(".env")) -> None:
    """Small fallback so .env works even before python-dotenv is installed."""
    if load_dotenv or not path.exists():
        return
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


load_env_file_fallback()


HOST = "0.0.0.0"
PORT = 2055
DATA_DIR = Path("data")
DB_PATH = DATA_DIR / "netflow.db"
EXCEL_PATH = DATA_DIR / "reporte_trafico.xlsx"
DEVICES_CSV = Path("devices.csv")
ENV_EXAMPLE = Path(".env.example")
DEFAULT_INTERNAL_NETWORKS = "192.168.1.0/24"
REPORT_INTERVAL_SECONDS = 60
STATUS_INTERVAL_SECONDS = 30
NO_PACKET_WARNING_SECONDS = 60
MIKROTIK_SYNC_INTERVAL_SECONDS = 5 * 60
MAX_FLOWS_IN_EXCEL = 10000
RETRY_PACKET_TIMEOUT_SECONDS = 60 * 60
DNS_RETRY_SECONDS = 24 * 60 * 60

SERVICE_PORTS = {
    80: "HTTP",
    443: "HTTPS",
    53: "DNS",
    25: "SMTP",
    587: "SMTP TLS",
    993: "IMAPS",
    995: "POP3S",
    3389: "RDP",
    22: "SSH",
    21: "FTP",
    123: "NTP",
}

DEVICE_CSV_EXAMPLE = """ip,nombre,area
192.168.1.10,PC Administración,Administración
192.168.1.20,PC Gerencia,Gerencia
192.168.1.30,Notebook Soporte,Sistemas
"""

ENV_EXAMPLE_CONTENT = """MIKROTIK_HOST=192.168.1.1
MIKROTIK_PORT=8728
MIKROTIK_USER=usuario_api
MIKROTIK_PASSWORD=clave_api
MIKROTIK_USE_SSL=false
INTERNAL_NETWORKS=192.168.1.0/24,192.168.88.0/24
"""


@dataclass
class Stats:
    packets_received: int = 0
    flows_saved: int = 0
    parse_errors: int = 0
    template_waiting: int = 0
    mikrotik_syncs: int = 0
    mikrotik_errors: int = 0
    last_packet_at: float | None = None
    last_report_at: float = 0
    last_status_at: float = 0
    last_no_packet_warning_at: float = 0
    last_mikrotik_sync_at: float = 0
    mikrotik_sync_running: bool = False
    last_flows: deque[str] = field(default_factory=lambda: deque(maxlen=5))


def now_text() -> str:
    return datetime.now().isoformat(timespec="seconds")


def setup_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


def normalize_spaces(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def parse_internal_networks() -> list[ipaddress._BaseNetwork]:
    value = os.getenv("INTERNAL_NETWORKS", DEFAULT_INTERNAL_NETWORKS)
    networks: list[ipaddress._BaseNetwork] = []
    for raw_item in value.split(","):
        item = raw_item.strip()
        if not item:
            continue
        try:
            networks.append(ipaddress.ip_network(item, strict=False))
        except ValueError:
            logging.warning("Red interna invalida en INTERNAL_NETWORKS: %s", item)
    return networks or [ipaddress.ip_network(DEFAULT_INTERNAL_NETWORKS)]


def is_truthy(value: str | None) -> bool:
    return str(value or "").strip().lower() in {"1", "true", "yes", "si", "y"}


def ensure_files() -> None:
    DATA_DIR.mkdir(exist_ok=True)
    if not DEVICES_CSV.exists():
        DEVICES_CSV.write_text(DEVICE_CSV_EXAMPLE, encoding="utf-8")
        logging.info("Creado devices.csv de ejemplo")
    if not ENV_EXAMPLE.exists():
        ENV_EXAMPLE.write_text(ENV_EXAMPLE_CONTENT, encoding="utf-8")
        logging.info("Creado .env.example")


def column_exists(conn: sqlite3.Connection, table: str, column: str) -> bool:
    rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
    return any(row[1] == column for row in rows)


def add_column_if_missing(
    conn: sqlite3.Connection, table: str, column: str, definition: str
) -> None:
    if not column_exists(conn, table, column):
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def init_db() -> None:
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS flows (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                received_at TEXT NOT NULL,
                exporter_ip TEXT,
                src_ip TEXT NOT NULL,
                dst_ip TEXT NOT NULL,
                src_port INTEGER,
                dst_port INTEGER,
                protocol INTEGER,
                bytes INTEGER NOT NULL DEFAULT 0,
                packets INTEGER NOT NULL DEFAULT 0,
                service TEXT NOT NULL,
                raw_json TEXT
            )
            """
        )
        add_column_if_missing(conn, "flows", "src_device_name", "TEXT")
        add_column_if_missing(conn, "flows", "dst_device_name", "TEXT")
        add_column_if_missing(conn, "flows", "src_mac", "TEXT")
        add_column_if_missing(conn, "flows", "dst_mac", "TEXT")

        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_flows_received_at ON flows(received_at)"
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_flows_src_ip ON flows(src_ip)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_flows_dst_ip ON flows(dst_ip)")

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS devices (
                ip TEXT PRIMARY KEY,
                mac TEXT,
                nombre TEXT,
                comment TEXT,
                active_host_name TEXT,
                host_name TEXT,
                source TEXT,
                area TEXT,
                last_seen DATETIME,
                updated_at DATETIME
            )
            """
        )
        add_column_if_missing(conn, "devices", "area", "TEXT")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_devices_source ON devices(source)")

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS dns_cache (
                ip TEXT PRIMARY KEY,
                hostname TEXT,
                resolved_at TEXT,
                error TEXT
            )
            """
        )
        conn.commit()


def normalize_ip(value: Any) -> str | None:
    if value is None:
        return None
    try:
        return str(ipaddress.ip_address(str(value).strip()))
    except ValueError:
        return None


def upsert_manual_device(conn: sqlite3.Connection, ip: str, nombre: str, area: str) -> None:
    current = conn.execute("SELECT * FROM devices WHERE ip = ?", (ip,)).fetchone()
    updated_at = now_text()
    if current and not nombre:
        conn.execute(
            """
            UPDATE devices
            SET area = COALESCE(NULLIF(?, ''), area), updated_at = ?
            WHERE ip = ?
            """,
            (area, updated_at, ip),
        )
        return

    if current:
        conn.execute(
            """
            UPDATE devices
            SET nombre = ?, area = ?, source = 'manual_csv', updated_at = ?
            WHERE ip = ?
            """,
            (nombre or ip, area, updated_at, ip),
        )
        return

    conn.execute(
        """
        INSERT INTO devices (
            ip, mac, nombre, comment, active_host_name, host_name,
            source, area, last_seen, updated_at
        )
        VALUES (?, '', ?, '', '', '', ?, ?, ?, ?)
        """,
        (ip, nombre or ip, "manual_csv" if nombre else "unknown", area, updated_at, updated_at),
    )


def sync_manual_devices_from_csv() -> int:
    if not DEVICES_CSV.exists():
        return 0

    count = 0
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        with DEVICES_CSV.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                ip = normalize_ip(row.get("ip"))
                if not ip:
                    continue
                nombre = normalize_spaces(row.get("nombre"))
                area = normalize_spaces(row.get("area"))
                upsert_manual_device(conn, ip, nombre, area)
                count += 1
        conn.commit()

    logging.info("Dispositivos cargados desde devices.csv: %s", count)
    return count


def choose_mikrotik_name(lease: dict[str, Any], ip: str, mac: str) -> tuple[str, str]:
    comment = normalize_spaces(lease.get("comment"))
    active_host_name = normalize_spaces(lease.get("active-host-name"))
    host_name = normalize_spaces(lease.get("host-name"))

    if comment:
        return comment, "mikrotik_comment"
    if active_host_name:
        return active_host_name, "mikrotik_active_host_name"
    if host_name:
        return host_name, "mikrotik_host_name"
    if mac:
        return mac, "mikrotik_mac"
    return ip, "unknown"


def upsert_mikrotik_device(conn: sqlite3.Connection, lease: dict[str, Any]) -> bool:
    ip = normalize_ip(lease.get("active-address") or lease.get("address"))
    if not ip:
        return False

    mac = normalize_spaces(lease.get("active-mac-address") or lease.get("mac-address"))
    comment = normalize_spaces(lease.get("comment"))
    active_host_name = normalize_spaces(lease.get("active-host-name"))
    host_name = normalize_spaces(lease.get("host-name"))
    nombre, source = choose_mikrotik_name(lease, ip, mac)
    now = now_text()

    current = conn.execute("SELECT * FROM devices WHERE ip = ?", (ip,)).fetchone()
    has_manual_name = (
        current
        and current["source"] == "manual_csv"
        and normalize_spaces(current["nombre"])
        and normalize_spaces(current["nombre"]) != ip
    )

    if has_manual_name:
        conn.execute(
            """
            UPDATE devices
            SET mac = COALESCE(NULLIF(?, ''), mac),
                comment = ?,
                active_host_name = ?,
                host_name = ?,
                last_seen = ?,
                updated_at = ?
            WHERE ip = ?
            """,
            (mac, comment, active_host_name, host_name, now, now, ip),
        )
        return True

    conn.execute(
        """
        INSERT INTO devices (
            ip, mac, nombre, comment, active_host_name, host_name,
            source, area, last_seen, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, COALESCE((SELECT area FROM devices WHERE ip = ?), ''), ?, ?)
        ON CONFLICT(ip) DO UPDATE SET
            mac = excluded.mac,
            nombre = excluded.nombre,
            comment = excluded.comment,
            active_host_name = excluded.active_host_name,
            host_name = excluded.host_name,
            source = excluded.source,
            last_seen = excluded.last_seen,
            updated_at = excluded.updated_at
        """,
        (ip, mac, nombre, comment, active_host_name, host_name, source, ip, now, now),
    )
    return True


def mikrotik_settings() -> dict[str, Any] | None:
    host = normalize_spaces(os.getenv("MIKROTIK_HOST"))
    user = normalize_spaces(os.getenv("MIKROTIK_USER"))
    password = os.getenv("MIKROTIK_PASSWORD", "")
    if not host or not user or not password:
        missing = []
        if not host:
            missing.append("MIKROTIK_HOST")
        if not user:
            missing.append("MIKROTIK_USER")
        if not password:
            missing.append("MIKROTIK_PASSWORD")
        logging.info(
            "MikroTik API no configurada; faltan variables: %s",
            ", ".join(missing),
        )
        return None

    return {
        "host": host,
        "username": user,
        "password": password,
        "port": int(os.getenv("MIKROTIK_PORT", "8728")),
        "use_ssl": is_truthy(os.getenv("MIKROTIK_USE_SSL")),
    }


def sync_mikrotik_devices() -> int:
    settings = mikrotik_settings()
    if not settings:
        logging.info("Se usara devices.csv si existe")
        return 0
    if routeros_connect is None:
        logging.warning("librouteros no esta instalado; no puedo sincronizar DHCP leases")
        return 0

    api = None
    try:
        kwargs = {
            "username": settings["username"],
            "password": settings["password"],
            "host": settings["host"],
            "port": settings["port"],
            "timeout": 5,
        }
        if settings["use_ssl"]:
            ctx = ssl.create_default_context()
            ctx.check_hostname = False
            ctx.verify_mode = ssl.CERT_NONE
            kwargs["ssl_wrapper"] = ctx.wrap_socket

        api = routeros_connect(**kwargs)
        leases = list(api.path("ip", "dhcp-server", "lease"))

        count = 0
        with sqlite3.connect(DB_PATH) as conn:
            conn.row_factory = sqlite3.Row
            for lease in leases:
                if upsert_mikrotik_device(conn, dict(lease)):
                    count += 1
            conn.commit()

        logging.info("Dispositivos sincronizados desde MikroTik DHCP leases: %s", count)
        return count
    except Exception as exc:
        logging.warning("No pude sincronizar dispositivos desde MikroTik: %s", exc)
        return 0
    finally:
        if api is not None:
            try:
                api.close()
            except Exception:
                pass


def load_device_lookup() -> dict[str, dict[str, str]]:
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            """
            SELECT ip, mac, nombre, comment, active_host_name, host_name,
                   source, area, last_seen, updated_at
            FROM devices
            """
        ).fetchall()
    return {row["ip"]: dict(row) for row in rows}


def is_internal_ip(ip: str, networks: list[ipaddress._BaseNetwork]) -> bool:
    try:
        address = ipaddress.ip_address(ip)
    except ValueError:
        return False
    return any(address.version == network.version and address in network for network in networks)


def device_name(
    ip: str,
    devices: dict[str, dict[str, str]],
    dns_names: dict[str, str] | None = None,
    internal_networks: list[ipaddress._BaseNetwork] | None = None,
) -> str:
    if ip in devices and normalize_spaces(devices[ip].get("nombre")):
        return normalize_spaces(devices[ip].get("nombre"))
    if internal_networks and is_internal_ip(ip, internal_networks):
        return ip
    if dns_names and dns_names.get(ip):
        return dns_names[ip]
    return ip


def device_area(ip: str, devices: dict[str, dict[str, str]]) -> str:
    return normalize_spaces(devices.get(ip, {}).get("area"))


def device_mac(ip: str, devices: dict[str, dict[str, str]]) -> str:
    return normalize_spaces(devices.get(ip, {}).get("mac"))


class ReverseDNSCache:
    """Resolve reverse DNS in the background so UDP collection is never blocked."""

    def __init__(self, db_path: Path, max_workers: int = 2) -> None:
        self.db_path = db_path
        self.executor = ThreadPoolExecutor(max_workers=max_workers)
        self.pending: set[str] = set()

    def has_recent_attempt(self, ip: str) -> bool:
        with sqlite3.connect(self.db_path) as conn:
            row = conn.execute(
                "SELECT resolved_at FROM dns_cache WHERE ip = ?",
                (ip,),
            ).fetchone()

        if not row or not row[0]:
            return False

        try:
            resolved_at = datetime.fromisoformat(row[0])
        except ValueError:
            return False
        return time.time() - resolved_at.timestamp() < DNS_RETRY_SECONDS

    def request(self, ip: str) -> None:
        if ip in self.pending or self.has_recent_attempt(ip):
            return
        self.pending.add(ip)
        future = self.executor.submit(self._resolve, ip)
        future.add_done_callback(lambda _: self.pending.discard(ip))

    def _resolve(self, ip: str) -> None:
        hostname = None
        error = None
        try:
            socket.setdefaulttimeout(1.0)
            hostname = socket.gethostbyaddr(ip)[0]
        except Exception as exc:  # DNS failures are expected and harmless.
            error = str(exc)[:200]

        with sqlite3.connect(self.db_path) as conn:
            conn.execute(
                """
                INSERT INTO dns_cache (ip, hostname, resolved_at, error)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(ip) DO UPDATE SET
                    hostname = excluded.hostname,
                    resolved_at = excluded.resolved_at,
                    error = excluded.error
                """,
                (ip, hostname, now_text(), error),
            )
            conn.commit()


def get_field(flow: Any, names: list[str], default: Any = None) -> Any:
    data = getattr(flow, "data", None)
    if isinstance(data, dict):
        for name in names:
            if name in data:
                return data[name]

    for name in names:
        if hasattr(flow, name):
            return getattr(flow, name)
    return default


def int_or_none(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def int_or_zero(value: Any) -> int:
    parsed = int_or_none(value)
    return parsed if parsed is not None else 0


def is_noise(
    src_ip: str, dst_ip: str, internal_networks: list[ipaddress._BaseNetwork]
) -> bool:
    src = ipaddress.ip_address(src_ip)
    dst = ipaddress.ip_address(dst_ip)

    if src.is_multicast or dst.is_multicast:
        return True
    if src.version == 4 and src == ipaddress.ip_address("255.255.255.255"):
        return True
    if dst.version == 4 and dst == ipaddress.ip_address("255.255.255.255"):
        return True
    for network in internal_networks:
        if src.version != network.version or dst.version != network.version:
            continue
        if src == network.broadcast_address or dst == network.broadcast_address:
            return True
        if src in network and dst in network:
            return True
    return False


def service_name(src_port: int | None, dst_port: int | None) -> str:
    if dst_port in SERVICE_PORTS:
        return SERVICE_PORTS[dst_port]
    if src_port in SERVICE_PORTS:
        return SERVICE_PORTS[src_port]
    return "Otro"


def extract_flow(
    flow: Any,
    received_at: str,
    exporter_ip: str,
    devices: dict[str, dict[str, str]],
    internal_networks: list[ipaddress._BaseNetwork],
) -> dict[str, Any] | None:
    src_ip = normalize_ip(
        get_field(flow, ["IPV4_SRC_ADDR", "IPV6_SRC_ADDR", "SRCADDR", "sourceIPv4Address"])
    )
    dst_ip = normalize_ip(
        get_field(flow, ["IPV4_DST_ADDR", "IPV6_DST_ADDR", "DSTADDR", "destinationIPv4Address"])
    )
    if not src_ip or not dst_ip or is_noise(src_ip, dst_ip, internal_networks):
        return None

    src_port = int_or_none(get_field(flow, ["L4_SRC_PORT", "SRCPORT", "sourceTransportPort"]))
    dst_port = int_or_none(get_field(flow, ["L4_DST_PORT", "DSTPORT", "destinationTransportPort"]))
    protocol = int_or_none(get_field(flow, ["PROTOCOL", "PROTO", "protocolIdentifier"]))
    bytes_count = int_or_zero(
        get_field(flow, ["IN_BYTES", "OUT_BYTES", "IN_PERMANENT_BYTES", "octetDeltaCount"])
    )
    packets_count = int_or_zero(
        get_field(flow, ["IN_PKTS", "OUT_PKTS", "IN_PERMANENT_PKTS", "packetDeltaCount"])
    )
    raw_data = getattr(flow, "data", None)
    if not isinstance(raw_data, dict):
        raw_data = dict(getattr(flow, "__dict__", {}))

    return {
        "received_at": received_at,
        "exporter_ip": exporter_ip,
        "src_ip": src_ip,
        "dst_ip": dst_ip,
        "src_port": src_port,
        "dst_port": dst_port,
        "protocol": protocol,
        "bytes": bytes_count,
        "packets": packets_count,
        "service": service_name(src_port, dst_port),
        "src_device_name": device_name(src_ip, devices, internal_networks=internal_networks),
        "dst_device_name": device_name(dst_ip, devices, internal_networks=internal_networks),
        "src_mac": device_mac(src_ip, devices),
        "dst_mac": device_mac(dst_ip, devices),
        "raw_json": json.dumps(raw_data, default=str, ensure_ascii=False),
    }


def save_flow(conn: sqlite3.Connection, item: dict[str, Any]) -> None:
    conn.execute(
        """
        INSERT INTO flows (
            received_at, exporter_ip, src_ip, dst_ip, src_port, dst_port,
            protocol, bytes, packets, service, src_device_name, dst_device_name,
            src_mac, dst_mac, raw_json
        )
        VALUES (
            :received_at, :exporter_ip, :src_ip, :dst_ip, :src_port, :dst_port,
            :protocol, :bytes, :packets, :service, :src_device_name, :dst_device_name,
            :src_mac, :dst_mac, :raw_json
        )
        """,
        item,
    )


def process_export(
    export: Any,
    received_at: str,
    exporter_ip: str,
    stats: Stats,
    dns_cache: ReverseDNSCache,
    internal_networks: list[ipaddress._BaseNetwork],
) -> int:
    saved = 0
    devices = load_device_lookup()
    with sqlite3.connect(DB_PATH) as conn:
        for flow in getattr(export, "flows", []):
            item = extract_flow(flow, received_at, exporter_ip, devices, internal_networks)
            if item is None:
                continue
            save_flow(conn, item)
            stats.last_flows.append(
                f"{item['src_device_name']} -> {item['dst_device_name']} "
                f"{item['service']} {item['bytes']} bytes"
            )
            if not is_internal_ip(item["dst_ip"], internal_networks):
                dns_cache.request(item["dst_ip"])
            saved += 1
        conn.commit()

    stats.flows_saved += saved
    return saved


def parse_payload(payload: bytes, templates: dict[str, dict[Any, Any]]) -> Any:
    return netflow.parse_packet(payload, templates)


def retry_waiting_packets(
    retry_packets: list[tuple[float, tuple[str, int], bytes]],
    templates: dict[str, dict[Any, Any]],
    stats: Stats,
    dns_cache: ReverseDNSCache,
    internal_networks: list[ipaddress._BaseNetwork],
) -> list[tuple[float, tuple[str, int], bytes]]:
    still_waiting = []
    for packet_ts, client, payload in retry_packets:
        if time.time() - packet_ts > RETRY_PACKET_TIMEOUT_SECONDS:
            logging.warning("Descartado paquete antiguo sin template reconocido")
            continue
        try:
            export = parse_payload(payload, templates)
            received_at = datetime.fromtimestamp(packet_ts).isoformat(timespec="seconds")
            process_export(export, received_at, client[0], stats, dns_cache, internal_networks)
        except (V9TemplateNotRecognized, IPFIXTemplateNotRecognized):
            still_waiting.append((packet_ts, client, payload))
        except Exception as exc:
            stats.parse_errors += 1
            logging.exception("Error reintentando paquete en espera: %s", exc)

    stats.template_waiting = len(still_waiting)
    return still_waiting


def human_bytes(value: int | float | None) -> str:
    size = float(value or 0)
    for unit in ["B", "KB", "MB", "GB", "TB"]:
        if size < 1024 or unit == "TB":
            return f"{size:.1f} {unit}" if unit != "B" else f"{int(size)} B"
        size /= 1024
    return f"{size:.1f} TB"


def mb(value: int | float | None) -> float:
    return round(float(value or 0) / (1024 * 1024), 2)


def autosize(ws) -> None:
    for column in ws.columns:
        width = 10
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(len(value) + 2, 55))
        ws.column_dimensions[column_letter].width = width


def style_table(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    autosize(ws)


def append_rows(ws, headers: list[str], rows: list[tuple[Any, ...]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_table(ws)


def fetch_rows(query: str, params: tuple[Any, ...] = ()) -> list[sqlite3.Row]:
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        return conn.execute(query, params).fetchall()


def fetch_one(query: str, params: tuple[Any, ...] = ()) -> sqlite3.Row:
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        return conn.execute(query, params).fetchone()


def fetch_all_flows() -> list[sqlite3.Row]:
    return fetch_rows(
        """
        SELECT *
        FROM flows
        ORDER BY received_at DESC, id DESC
        LIMIT ?
        """,
        (MAX_FLOWS_IN_EXCEL,),
    )


def fetch_summary() -> sqlite3.Row:
    return fetch_one(
        """
        SELECT
            COUNT(*) AS total_flows,
            COALESCE(SUM(bytes), 0) AS total_bytes,
            COALESCE(SUM(packets), 0) AS total_packets,
            MIN(received_at) AS first_seen,
            MAX(received_at) AS last_seen
        FROM flows
        """
    )


def load_dns_cache() -> dict[str, str]:
    rows = fetch_rows("SELECT ip, hostname FROM dns_cache WHERE hostname IS NOT NULL")
    return {row["ip"]: row["hostname"] for row in rows}


def principal_destination(ip: str, devices: dict[str, dict[str, str]], dns_names: dict[str, str]) -> str:
    row = fetch_one(
        """
        SELECT peer_ip, SUM(bytes) AS total_bytes
        FROM (
            SELECT dst_ip AS peer_ip, bytes FROM flows WHERE src_ip = ?
            UNION ALL
            SELECT src_ip AS peer_ip, bytes FROM flows WHERE dst_ip = ?
        )
        GROUP BY peer_ip
        ORDER BY total_bytes DESC
        LIMIT 1
        """,
        (ip, ip),
    )
    if not row:
        return ""
    return device_name(row["peer_ip"], devices, dns_names=dns_names)


def principal_service(ip: str) -> str:
    row = fetch_one(
        """
        SELECT service, SUM(bytes) AS total_bytes
        FROM flows
        WHERE src_ip = ? OR dst_ip = ?
        GROUP BY service
        ORDER BY total_bytes DESC
        LIMIT 1
        """,
        (ip, ip),
    )
    return row["service"] if row else ""


def device_traffic_rows(
    devices: dict[str, dict[str, str]], dns_names: dict[str, str], internal_networks: list[ipaddress._BaseNetwork]
) -> list[tuple[Any, ...]]:
    ips_from_flows = {
        row["ip"]
        for row in fetch_rows(
            """
            SELECT src_ip AS ip FROM flows
            UNION
            SELECT dst_ip AS ip FROM flows
            """
        )
        if is_internal_ip(row["ip"], internal_networks)
    }
    ips_from_devices = {
        ip for ip in devices if is_internal_ip(ip, internal_networks)
    }
    rows = []
    for ip in sorted(ips_from_flows | ips_from_devices):
        sent = fetch_one(
            "SELECT COALESCE(SUM(bytes), 0) AS bytes, COUNT(*) AS flows, MAX(received_at) AS last_seen FROM flows WHERE src_ip = ?",
            (ip,),
        )
        received = fetch_one(
            "SELECT COALESCE(SUM(bytes), 0) AS bytes, COUNT(*) AS flows, MAX(received_at) AS last_seen FROM flows WHERE dst_ip = ?",
            (ip,),
        )
        sent_bytes = sent["bytes"] if sent else 0
        received_bytes = received["bytes"] if received else 0
        flow_count = (sent["flows"] if sent else 0) + (received["flows"] if received else 0)
        last_activity = max(
            [value for value in [sent["last_seen"] if sent else None, received["last_seen"] if received else None] if value],
            default="",
        )
        rows.append(
            (
                device_name(ip, devices, dns_names=dns_names, internal_networks=internal_networks),
                ip,
                device_mac(ip, devices),
                device_area(ip, devices),
                mb(sent_bytes),
                mb(received_bytes),
                mb(sent_bytes + received_bytes),
                flow_count,
                principal_destination(ip, devices, dns_names),
                principal_service(ip),
                last_activity,
            )
        )
    rows.sort(key=lambda item: item[6], reverse=True)
    return rows


def generate_excel(stats: Stats, internal_networks: list[ipaddress._BaseNetwork]) -> None:
    devices = load_device_lookup()
    dns_names = load_dns_cache()
    flows = fetch_all_flows()
    summary = fetch_summary()
    device_rows = fetch_rows(
        """
        SELECT src_ip, MAX(src_mac) AS src_mac, COALESCE(SUM(bytes), 0) AS bytes,
               COALESCE(SUM(packets), 0) AS packets, COUNT(*) AS flows
        FROM flows
        GROUP BY src_ip
        ORDER BY bytes DESC
        """
    )
    destination_rows = fetch_rows(
        """
        SELECT dst_ip, COALESCE(SUM(bytes), 0) AS bytes,
               COALESCE(SUM(packets), 0) AS packets, COUNT(*) AS flows
        FROM flows
        GROUP BY dst_ip
        ORDER BY bytes DESC
        """
    )
    service_rows = fetch_rows(
        """
        SELECT service, COALESCE(SUM(bytes), 0) AS bytes,
               COALESCE(SUM(packets), 0) AS packets, COUNT(*) AS flows
        FROM flows
        GROUP BY service
        ORDER BY bytes DESC
        """
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    append_rows(
        ws,
        ["Metrica", "Valor"],
        [
            ("Generado", now_text()),
            ("Redes internas", ", ".join(str(network) for network in internal_networks)),
            ("Flows en base", summary["total_flows"]),
            ("Bytes totales", summary["total_bytes"]),
            ("Bytes legibles", human_bytes(summary["total_bytes"])),
            ("Packets totales", summary["total_packets"]),
            ("Primer flow", summary["first_seen"] or ""),
            ("Ultimo flow", summary["last_seen"] or ""),
            ("Paquetes UDP recibidos", stats.packets_received),
            ("Flows guardados esta ejecucion", stats.flows_saved),
            ("Errores de parseo esta ejecucion", stats.parse_errors),
            ("Paquetes esperando template", stats.template_waiting),
            ("Sincronizaciones MikroTik", stats.mikrotik_syncs),
            ("Errores MikroTik", stats.mikrotik_errors),
        ],
    )

    ws = wb.create_sheet("Flows")
    flow_rows = []
    for row in flows:
        src_name = row["src_device_name"] or device_name(
            row["src_ip"], devices, dns_names=dns_names, internal_networks=internal_networks
        )
        dst_name = row["dst_device_name"] or device_name(
            row["dst_ip"], devices, dns_names=dns_names, internal_networks=internal_networks
        )
        if not is_internal_ip(row["dst_ip"], internal_networks) and dns_names.get(row["dst_ip"]):
            dst_name = dns_names[row["dst_ip"]]
        flow_rows.append(
            (
                src_name,
                dst_name,
                row["src_ip"],
                row["dst_ip"],
                row["service"],
                row["bytes"],
                mb(row["bytes"]),
                row["received_at"],
                row["src_port"],
                row["dst_port"],
                row["protocol"],
                row["packets"],
                row["src_mac"] or device_mac(row["src_ip"], devices),
                row["dst_mac"] or device_mac(row["dst_ip"], devices),
            )
        )
    append_rows(
        ws,
        [
            "Origen Nombre",
            "Destino Nombre",
            "Origen IP",
            "Destino IP",
            "Servicio",
            "Bytes",
            "MB",
            "Fecha",
            "Puerto origen",
            "Puerto destino",
            "Protocolo",
            "Packets",
            "Origen MAC",
            "Destino MAC",
        ],
        flow_rows,
    )

    ws = wb.create_sheet("Tráfico por dispositivo")
    append_rows(
        ws,
        [
            "Dispositivo",
            "IP",
            "MAC",
            "Área",
            "Total MB enviados",
            "Total MB recibidos",
            "Total MB total",
            "Flows",
            "Principal destino",
            "Principal servicio",
            "Ultima actividad",
        ],
        device_traffic_rows(devices, dns_names, internal_networks),
    )

    ws = wb.create_sheet("Top dispositivos")
    append_rows(
        ws,
        ["Dispositivo", "IP", "MAC", "Área", "Bytes", "MB", "Bytes legibles", "Packets", "Flows"],
        [
            (
                device_name(row["src_ip"], devices, dns_names=dns_names, internal_networks=internal_networks),
                row["src_ip"],
                row["src_mac"] or device_mac(row["src_ip"], devices),
                device_area(row["src_ip"], devices),
                row["bytes"],
                mb(row["bytes"]),
                human_bytes(row["bytes"]),
                row["packets"],
                row["flows"],
            )
            for row in device_rows
        ],
    )

    ws = wb.create_sheet("Top destinos")
    append_rows(
        ws,
        ["Destino", "IP destino", "DNS destino", "Bytes", "MB", "Bytes legibles", "Packets", "Flows"],
        [
            (
                device_name(row["dst_ip"], devices, dns_names=dns_names, internal_networks=internal_networks),
                row["dst_ip"],
                dns_names.get(row["dst_ip"], ""),
                row["bytes"],
                mb(row["bytes"]),
                human_bytes(row["bytes"]),
                row["packets"],
                row["flows"],
            )
            for row in destination_rows
        ],
    )

    ws = wb.create_sheet("Top servicios")
    append_rows(
        ws,
        ["Servicio", "Bytes", "MB", "Bytes legibles", "Packets", "Flows"],
        [
            (
                row["service"],
                row["bytes"],
                mb(row["bytes"]),
                human_bytes(row["bytes"]),
                row["packets"],
                row["flows"],
            )
            for row in service_rows
        ],
    )

    ws = wb.create_sheet("Dispositivos MikroTik")
    device_inventory_rows = fetch_rows(
        """
        SELECT ip, mac, nombre, comment, active_host_name, host_name,
               source, last_seen, updated_at
        FROM devices
        ORDER BY ip
        """
    )
    append_rows(
        ws,
        [
            "IP",
            "MAC",
            "Nombre visible",
            "Comment MikroTik",
            "Active Host Name",
            "Host Name",
            "Source",
            "Ultima vez visto",
            "Actualizado",
        ],
        [
            (
                row["ip"],
                row["mac"],
                row["nombre"],
                row["comment"],
                row["active_host_name"],
                row["host_name"],
                row["source"],
                row["last_seen"],
                row["updated_at"],
            )
            for row in device_inventory_rows
        ],
    )

    wb.save(EXCEL_PATH)
    logging.info("Excel actualizado: %s", EXCEL_PATH)


def log_status(stats: Stats) -> None:
    logging.info(
        "Estado: %s paquetes UDP, %s flows guardados, %s errores de parseo, %s esperando template",
        stats.packets_received,
        stats.flows_saved,
        stats.parse_errors,
        stats.template_waiting,
    )
    logging.info(
        "MikroTik: %s sincronizaciones, %s errores",
        stats.mikrotik_syncs,
        stats.mikrotik_errors,
    )
    if stats.last_flows:
        logging.info("Ultimos flows procesados: %s", " | ".join(stats.last_flows))


def maybe_warn_no_packets(stats: Stats) -> None:
    now = time.time()
    no_packets_yet = stats.last_packet_at is None
    too_long_since_packet = stats.last_packet_at is not None and (
        now - stats.last_packet_at >= NO_PACKET_WARNING_SECONDS
    )
    can_warn = now - stats.last_no_packet_warning_at >= NO_PACKET_WARNING_SECONDS
    if (no_packets_yet or too_long_since_packet) and can_warn:
        logging.warning(
            "No llegan paquetes NetFlow. Verificar MikroTik, firewall de Windows, IP 192.168.1.114 y UDP %s.",
            PORT,
        )
        stats.last_no_packet_warning_at = now


def schedule_mikrotik_sync(stats: Stats, executor: ThreadPoolExecutor) -> None:
    if stats.mikrotik_sync_running:
        return

    stats.mikrotik_sync_running = True

    def task() -> int:
        sync_manual_devices_from_csv()
        return sync_mikrotik_devices()

    future = executor.submit(task)

    def done_callback(done_future) -> None:
        stats.mikrotik_sync_running = False
        stats.last_mikrotik_sync_at = time.time()
        try:
            done_future.result()
            stats.mikrotik_syncs += 1
        except Exception as exc:
            stats.mikrotik_errors += 1
            logging.warning("Error sincronizando dispositivos: %s", exc)

    future.add_done_callback(done_callback)


def run_collector(internal_networks: list[ipaddress._BaseNetwork]) -> None:
    stats = Stats()
    stats.last_mikrotik_sync_at = time.time()
    dns_cache = ReverseDNSCache(DB_PATH)
    sync_executor = ThreadPoolExecutor(max_workers=1)
    templates: dict[str, dict[Any, Any]] = {"netflow": {}, "ipfix": {}}
    retry_packets: list[tuple[float, tuple[str, int], bytes]] = []

    sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    sock.bind((HOST, PORT))
    sock.settimeout(1.0)

    logging.info("Escuchando NetFlow en %s:%s UDP", HOST, PORT)
    logging.info("Redes internas: %s", ", ".join(str(network) for network in internal_networks))
    logging.info("Base SQLite: %s", DB_PATH)
    logging.info("Excel: %s", EXCEL_PATH)

    try:
        while True:
            now = time.time()

            try:
                payload, client = sock.recvfrom(65535)
            except socket.timeout:
                maybe_warn_no_packets(stats)
            except OSError as exc:
                stats.parse_errors += 1
                logging.exception("Error de socket: %s", exc)
                time.sleep(1)
            else:
                stats.packets_received += 1
                stats.last_packet_at = now
                received_at = now_text()

                try:
                    export = parse_payload(payload, templates)
                    saved = process_export(
                        export, received_at, client[0], stats, dns_cache, internal_networks
                    )
                    logging.debug("Procesados %s flows desde %s", saved, client[0])

                    if (
                        getattr(export, "header", None)
                        and export.header.version in [9, 10]
                        and getattr(export, "contains_new_templates", False)
                        and retry_packets
                    ):
                        retry_packets = retry_waiting_packets(
                            retry_packets, templates, stats, dns_cache, internal_networks
                        )
                except (V9TemplateNotRecognized, IPFIXTemplateNotRecognized):
                    retry_packets.append((now, client, payload))
                    stats.template_waiting = len(retry_packets)
                    logging.info(
                        "Paquete v9/IPFIX esperando template. En espera: %s",
                        stats.template_waiting,
                    )
                except UnknownExportVersion as exc:
                    stats.parse_errors += 1
                    logging.warning("Version NetFlow no reconocida: %s", exc)
                except Exception as exc:
                    stats.parse_errors += 1
                    logging.exception("Error de parseo/proceso: %s", exc)

            if time.time() - stats.last_mikrotik_sync_at >= MIKROTIK_SYNC_INTERVAL_SECONDS:
                schedule_mikrotik_sync(stats, sync_executor)

            if time.time() - stats.last_report_at >= REPORT_INTERVAL_SECONDS:
                try:
                    generate_excel(stats, internal_networks)
                except PermissionError:
                    logging.warning("No pude actualizar el Excel. Cerralo si esta abierto: %s", EXCEL_PATH)
                except Exception as exc:
                    logging.exception("Error generando Excel: %s", exc)
                stats.last_report_at = time.time()

            if time.time() - stats.last_status_at >= STATUS_INTERVAL_SECONDS:
                log_status(stats)
                stats.last_status_at = time.time()
    except KeyboardInterrupt:
        logging.info("Collector detenido por el usuario")
    finally:
        sock.close()
        sync_executor.shutdown(wait=False, cancel_futures=True)


def main() -> None:
    setup_logging()
    internal_networks = parse_internal_networks()
    ensure_files()
    init_db()
    sync_manual_devices_from_csv()
    sync_mikrotik_devices()
    generate_excel(Stats(), internal_networks)
    run_collector(internal_networks)


if __name__ == "__main__":
    main()
